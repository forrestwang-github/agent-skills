from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from bill_analyzer import analyze_records
from bill_importer import BillFileImporter
from bill_report import SHEET_NAMES, generate_report
from bill_validator import validate_import
from test_importer_validator import create_fixture


class AnalyzerReportTests(unittest.TestCase):
    def test_aggregation_and_report(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            temp = Path(tmp)
            source = temp / "bill.xlsx"
            create_fixture(source)
            imported = BillFileImporter().load(source)
            validation = validate_import(
                imported, expected_months=["2026-01", "2026-02"]
            )
            self.assertTrue(validation.passed, validation.to_dict())
            analysis = analyze_records(
                validation,
                customer_name="测试客户",
                group_by_region=True,
                top_n=5,
            )
            data = analysis.to_dict()
            self.assertEqual("111.105", data["totals"]["payable_amount"])
            self.assertEqual("11.005", data["totals"]["postpaid_amount"])
            self.assertEqual("passed", data["verification"]["sqlite_cross_check"])
            self.assertEqual(2, len(data["postpaid_item_rows"]))

            output = temp / "report.xlsx"
            generate_report(data, imported.manifest, validation, output)
            self.assertTrue(output.exists())

            workbook = openpyxl.load_workbook(
                output, read_only=True, data_only=True
            )
            try:
                self.assertEqual(SHEET_NAMES, workbook.sheetnames)
                self.assertGreater(workbook["账单分析"].max_row, 5)
                self.assertGreater(workbook["后付费分析"].max_column, 10)
            finally:
                workbook.close()

    def test_units_are_not_mixed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "bill.xlsx"
            create_fixture(source)
            imported = BillFileImporter().load(source)
            validation = validate_import(imported)
            analysis = analyze_records(validation)
            units = {
                row["usage_unit"]
                for row in analysis.data["postpaid_item_rows"]
            }
            self.assertEqual({"GB", "次"}, units)


if __name__ == "__main__":
    unittest.main()
