#!/usr/bin/env python3
"""Generate and verify the four-sheet bill analysis workbook."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from bill_models import OTHER, PAY_AS_YOU_GO, SUBSCRIPTION, ValidationResult


SHEET_NAMES = ["账单分析", "后付费分析", "趋势分析", "数据说明"]
TYPE_ORDER = [SUBSCRIPTION, PAY_AS_YOU_GO, OTHER]


def generate_report(
    analysis: Dict[str, Any],
    manifest: Dict[str, Any],
    validation: ValidationResult,
    output_path: Path,
) -> Path:
    try:
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，无法生成 Excel") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    styles = _build_styles(Font, PatternFill, Border, Side, Alignment)

    sheet1 = workbook.create_sheet("账单分析")
    _build_analysis_sheet(sheet1, analysis, styles, get_column_letter)

    sheet2 = workbook.create_sheet("后付费分析")
    _build_postpaid_sheet(sheet2, analysis, styles, get_column_letter)

    sheet3 = workbook.create_sheet("趋势分析")
    _build_trend_sheet(
        sheet3,
        analysis,
        styles,
        get_column_letter,
        LineChart,
        Reference,
    )

    sheet4 = workbook.create_sheet("数据说明")
    _build_notes_sheet(
        sheet4,
        analysis,
        manifest,
        validation,
        styles,
        get_column_letter,
    )

    workbook.properties.creator = "aliyun-bill-analysis-skill"
    workbook.properties.title = "阿里云账单分析"
    workbook.properties.description = (
        "程序确定性生成的产品、计费方式、计费项、用量和费用趋势报告"
    )
    workbook.save(output_path)
    verify_report(output_path, analysis)
    return output_path


def _build_styles(Font, PatternFill, Border, Side, Alignment) -> Dict[str, Any]:
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return {
        "title_font": Font(name="微软雅黑", size=16, bold=True, color="FFFFFF"),
        "title_fill": PatternFill("solid", fgColor="1F4E78"),
        "header_font": Font(name="微软雅黑", size=10, bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor="4472C4"),
        "sub_fill": PatternFill("solid", fgColor="D9EAF7"),
        "total_fill": PatternFill("solid", fgColor="BDD7EE"),
        "warning_fill": PatternFill("solid", fgColor="FFF2CC"),
        "error_fill": PatternFill("solid", fgColor="F4CCCC"),
        "body_font": Font(name="微软雅黑", size=10),
        "border": border,
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "wrap": Alignment(
            horizontal="left", vertical="top", wrap_text=True
        ),
    }


def _set_title(
    worksheet,
    title: str,
    subtitle: str,
    last_column: int,
    styles: Dict[str, Any],
) -> None:
    worksheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=last_column
    )
    cell = worksheet.cell(1, 1, title)
    cell.font = styles["title_font"]
    cell.fill = styles["title_fill"]
    cell.alignment = styles["center"]
    worksheet.row_dimensions[1].height = 28
    worksheet.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=last_column
    )
    subtitle_cell = worksheet.cell(2, 1, subtitle)
    subtitle_cell.font = styles["body_font"]
    subtitle_cell.alignment = styles["left"]
    worksheet.row_dimensions[2].height = 22


def _style_header(
    worksheet, row: int, headers: List[str], styles: Dict[str, Any]
) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row, column, header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.border = styles["border"]
        cell.alignment = styles["center"]
    worksheet.row_dimensions[row].height = 30


def _style_body_row(
    worksheet,
    row: int,
    column_count: int,
    styles: Dict[str, Any],
    fill: Optional[Any] = None,
) -> None:
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row, column)
        cell.font = styles["body_font"]
        cell.border = styles["border"]
        cell.alignment = styles["left"] if column <= 3 else styles["center"]
        if fill:
            cell.fill = fill


def _build_analysis_sheet(
    worksheet,
    analysis: Dict[str, Any],
    styles: Dict[str, Any],
    get_column_letter,
) -> None:
    months = analysis["months"]
    headers = ["计费方式", "产品Code", "产品名称", *months, "合计"]
    last_column = len(headers)
    subtitle = _subtitle(analysis)
    _set_title(worksheet, "账单分析", subtitle, last_column, styles)
    _style_header(worksheet, 4, headers, styles)

    row_number = 5
    for billing_type in TYPE_ORDER:
        group = [
            row
            for row in analysis["product_rows"]
            if row["billing_type"] == billing_type
        ]
        if not group:
            continue
        type_month_totals = {month: Decimal("0") for month in months}
        type_total = Decimal("0")
        for row in group:
            values: List[Any] = [
                row["billing_type_display"],
                row["product_code"],
                row["product_name"],
            ]
            for month in months:
                amount = Decimal(row["monthly_amounts"][month])
                type_month_totals[month] += amount
                values.append(_excel_number(amount))
            row_total = Decimal(row["total_amount"])
            type_total += row_total
            values.append(_excel_number(row_total))
            for column, value in enumerate(values, start=1):
                worksheet.cell(row_number, column, value)
            _style_body_row(
                worksheet, row_number, last_column, styles
            )
            _format_amount_columns(
                worksheet, row_number, 4, last_column
            )
            row_number += 1

        subtotal_values = [
            f"{group[0]['billing_type_display']}小计",
            "",
            "",
            *[
                _excel_number(type_month_totals[month])
                for month in months
            ],
            _excel_number(type_total),
        ]
        for column, value in enumerate(subtotal_values, start=1):
            worksheet.cell(row_number, column, value)
        _style_body_row(
            worksheet,
            row_number,
            last_column,
            styles,
            fill=styles["sub_fill"],
        )
        worksheet.cell(row_number, 1).font = _bold(worksheet.cell(row_number, 1).font)
        _format_amount_columns(worksheet, row_number, 4, last_column)
        row_number += 1

    total = Decimal(analysis["totals"]["payable_amount"])
    grand_values = [
        "总合计",
        "",
        "",
        *[
            _excel_number(Decimal(analysis["monthly_total"][month]))
            for month in months
        ],
        _excel_number(total),
    ]
    for column, value in enumerate(grand_values, start=1):
        worksheet.cell(row_number, column, value)
    _style_body_row(
        worksheet,
        row_number,
        last_column,
        styles,
        fill=styles["total_fill"],
    )
    for column in range(1, last_column + 1):
        worksheet.cell(row_number, column).font = _bold(
            worksheet.cell(row_number, column).font
        )
    _format_amount_columns(worksheet, row_number, 4, last_column)

    worksheet.freeze_panes = "D5"
    worksheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{row_number}"
    widths = {1: 14, 2: 18, 3: 28}
    _set_widths(worksheet, widths, default_start=4, default_width=16)


def _build_postpaid_sheet(
    worksheet,
    analysis: Dict[str, Any],
    styles: Dict[str, Any],
    get_column_letter,
) -> None:
    months = analysis["months"]
    headers = [
        "产品Code",
        "产品名称",
        "计费项Code",
        "计费项名称",
        "地域Code",
        "地域",
        "用量单位",
    ]
    for month in months:
        headers.extend([f"{month}用量", f"{month}费用"])
    headers.extend(["总用量", "总费用"])
    last_column = len(headers)
    _set_title(
        worksheet,
        "后付费分析",
        _subtitle(analysis),
        last_column,
        styles,
    )
    _style_header(worksheet, 4, headers, styles)

    row_number = 5
    for row in analysis["postpaid_item_rows"]:
        values: List[Any] = [
            row["product_code"],
            row["product_name"],
            row["item_code"],
            row["item_name"],
            row["region_code"],
            row["region_name"],
            row["usage_unit"],
        ]
        for month in months:
            usage = row["monthly"][month]["usage"]
            amount = row["monthly"][month]["amount"]
            values.extend(
                [
                    _excel_number(Decimal(usage)) if usage is not None else None,
                    _excel_number(Decimal(amount)),
                ]
            )
        values.extend(
            [
                _excel_number(Decimal(row["total_usage"]))
                if row["total_usage"] is not None
                else None,
                _excel_number(Decimal(row["total_amount"])),
            ]
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_number, column, value)
        _style_body_row(worksheet, row_number, last_column, styles)
        for column in range(8, 8 + len(months) * 2, 2):
            worksheet.cell(row_number, column).number_format = "0.########"
            worksheet.cell(row_number, column + 1).number_format = "#,##0.00####"
        worksheet.cell(row_number, last_column - 1).number_format = "0.########"
        worksheet.cell(row_number, last_column).number_format = "#,##0.00####"
        row_number += 1

    total = Decimal(analysis["totals"]["postpaid_amount"])
    total_values: List[Any] = ["总计", "", "", "", "", "", ""]
    for month in months:
        total_values.extend(
            [
                None,
                _excel_number(
                    Decimal(
                        analysis["billing_type_monthly"][PAY_AS_YOU_GO][month]
                    )
                ),
            ]
        )
    total_values.extend([None, _excel_number(total)])
    for column, value in enumerate(total_values, start=1):
        worksheet.cell(row_number, column, value)
    _style_body_row(
        worksheet,
        row_number,
        last_column,
        styles,
        fill=styles["total_fill"],
    )
    for column in range(1, last_column + 1):
        worksheet.cell(row_number, column).font = _bold(
            worksheet.cell(row_number, column).font
        )
    for column in range(8, 8 + len(months) * 2, 2):
        worksheet.cell(row_number, column + 1).number_format = "#,##0.00####"
    worksheet.cell(row_number, last_column).number_format = "#,##0.00####"

    worksheet.freeze_panes = "H5"
    worksheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{row_number}"
    widths = {
        1: 18,
        2: 28,
        3: 22,
        4: 30,
        5: 16,
        6: 18,
        7: 15,
    }
    _set_widths(worksheet, widths, default_start=8, default_width=15)


def _build_trend_sheet(
    worksheet,
    analysis: Dict[str, Any],
    styles: Dict[str, Any],
    get_column_letter,
    LineChart,
    Reference,
) -> None:
    months = analysis["months"]
    top_n = max(1, int(analysis.get("top_n", 10)))
    _set_title(
        worksheet,
        "趋势分析",
        "趋势仅描述账单金额和计费用量变化，不代表资源利用率或业务原因。",
        max(12, len(months) + 8),
        styles,
    )

    monthly_headers = ["月份", "总费用", "包年包月", "后付费", "其他"]
    _style_header(worksheet, 4, monthly_headers, styles)
    for index, month in enumerate(months, start=5):
        values = [
            month,
            _excel_number(Decimal(analysis["monthly_total"][month])),
            _excel_number(
                Decimal(
                    analysis["billing_type_monthly"][SUBSCRIPTION][month]
                )
            ),
            _excel_number(
                Decimal(
                    analysis["billing_type_monthly"][PAY_AS_YOU_GO][month]
                )
            ),
            _excel_number(
                Decimal(analysis["billing_type_monthly"][OTHER][month])
            ),
        ]
        for column, value in enumerate(values, start=1):
            worksheet.cell(index, column, value)
        _style_body_row(worksheet, index, len(monthly_headers), styles)
        _format_amount_columns(worksheet, index, 2, 5)

    chart = LineChart()
    chart.title = "月度费用趋势"
    chart.y_axis.title = "费用"
    chart.x_axis.title = "月份"
    data = Reference(
        worksheet,
        min_col=2,
        max_col=5,
        min_row=4,
        max_row=4 + len(months),
    )
    categories = Reference(
        worksheet, min_col=1, min_row=5, max_row=4 + len(months)
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 15
    worksheet.add_chart(chart, "G4")

    top_products = sorted(
        _collapse_product_rows(analysis["product_rows"], months),
        key=lambda row: -Decimal(row["total_amount"]),
    )[:top_n]
    product_start = 7 + len(months)
    product_headers = ["产品Code", "产品名称", *months, "合计"]
    _style_header(worksheet, product_start, product_headers, styles)
    for offset, row in enumerate(top_products, start=1):
        excel_row = product_start + offset
        values = [
            row["product_code"],
            row["product_name"],
            *[
                _excel_number(Decimal(row["monthly_amounts"][month]))
                for month in months
            ],
            _excel_number(Decimal(row["total_amount"])),
        ]
        for column, value in enumerate(values, start=1):
            worksheet.cell(excel_row, column, value)
        _style_body_row(
            worksheet, excel_row, len(product_headers), styles
        )
        _format_amount_columns(
            worksheet, excel_row, 3, len(product_headers)
        )

    if top_products and months:
        product_chart = LineChart()
        product_chart.title = f"费用最高的前{len(top_products)}个产品"
        product_chart.y_axis.title = "费用"
        product_chart.x_axis.title = "月份"
        for row_index in range(
            product_start + 1, product_start + 1 + len(top_products)
        ):
            values = Reference(
                worksheet,
                min_col=3,
                max_col=2 + len(months),
                min_row=row_index,
                max_row=row_index,
            )
            series = openpyxl_chart_series(
                values,
                title=f"{worksheet.cell(row_index, 2).value}",
            )
            product_chart.series.append(series)
        product_categories = Reference(
            worksheet,
            min_col=3,
            max_col=2 + len(months),
            min_row=product_start,
            max_row=product_start,
        )
        product_chart.set_categories(product_categories)
        product_chart.height = 9
        product_chart.width = 18
        worksheet.add_chart(
            product_chart,
            f"{get_column_letter(len(product_headers) + 2)}{product_start}",
        )

    item_start = product_start + max(2, len(top_products) + 3)
    item_headers = [
        "产品",
        "计费项",
        "地域",
        "单位",
    ]
    for month in months:
        item_headers.extend([f"{month}用量", f"{month}费用"])
    item_headers.extend(["总用量", "总费用"])
    _style_header(worksheet, item_start, item_headers, styles)
    top_items = analysis["postpaid_item_rows"][:top_n]
    for offset, row in enumerate(top_items, start=1):
        excel_row = item_start + offset
        values: List[Any] = [
            row["product_name"],
            row["item_name"],
            row["region_name"] or row["region_code"],
            row["usage_unit"],
        ]
        for month in months:
            usage = row["monthly"][month]["usage"]
            amount = row["monthly"][month]["amount"]
            values.extend(
                [
                    _excel_number(Decimal(usage)) if usage is not None else None,
                    _excel_number(Decimal(amount)),
                ]
            )
        values.extend(
            [
                _excel_number(Decimal(row["total_usage"]))
                if row["total_usage"] is not None
                else None,
                _excel_number(Decimal(row["total_amount"])),
            ]
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(excel_row, column, value)
        _style_body_row(
            worksheet, excel_row, len(item_headers), styles
        )
        for column in range(5, 5 + len(months) * 2, 2):
            worksheet.cell(excel_row, column).number_format = "0.########"
            worksheet.cell(excel_row, column + 1).number_format = "#,##0.00####"
        worksheet.cell(excel_row, len(item_headers) - 1).number_format = "0.########"
        worksheet.cell(excel_row, len(item_headers)).number_format = "#,##0.00####"

    worksheet.freeze_panes = "A5"
    _set_widths(
        worksheet,
        {1: 18, 2: 28, 3: 18, 4: 14, 6: 18, 7: 18},
        default_start=5,
        default_width=15,
    )


def openpyxl_chart_series(values, title: str):
    from openpyxl.chart import Series

    return Series(values, title=title)


def _build_notes_sheet(
    worksheet,
    analysis: Dict[str, Any],
    manifest: Dict[str, Any],
    validation: ValidationResult,
    styles: Dict[str, Any],
    get_column_letter,
) -> None:
    _set_title(
        worksheet,
        "数据说明",
        "本页记录数据范围、解析结果、校验结论和异常，不包含AK/SK。",
        5,
        styles,
    )
    row_number = 4

    def section(title: str) -> None:
        nonlocal row_number
        worksheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=5,
        )
        cell = worksheet.cell(row_number, 1, title)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["left"]
        row_number += 1

    def key_value(key: str, value: Any) -> None:
        nonlocal row_number
        worksheet.cell(row_number, 1, key)
        worksheet.merge_cells(
            start_row=row_number,
            start_column=2,
            end_row=row_number,
            end_column=5,
        )
        worksheet.cell(row_number, 2, "" if value is None else str(value))
        _style_body_row(worksheet, row_number, 5, styles)
        worksheet.cell(row_number, 2).alignment = styles["wrap"]
        row_number += 1

    section("分析范围")
    key_value("客户名称", analysis.get("customer_name", ""))
    key_value("数据来源", manifest.get("source", ""))
    key_value("账期", " ~ ".join([analysis["months"][0], analysis["months"][-1]]))
    key_value("币种", ", ".join(analysis.get("currencies", [])) or "未提供")
    key_value("地域维度", "启用" if analysis.get("group_by_region") else "关闭")
    key_value("生成时间", datetime.now().astimezone().isoformat(timespec="seconds"))

    section("质量结论")
    key_value("校验状态", "通过" if validation.passed else "失败")
    for key, value in validation.stats.items():
        key_value(key, value)
    key_value("SQLite交叉校验", analysis["verification"]["sqlite_cross_check"])

    section("来源文件")
    _style_header(
        worksheet,
        row_number,
        ["文件", "SHA-256", "原始行数", "解析行数", "Sheet"],
        styles,
    )
    row_number += 1
    for file_info in manifest.get("files", []):
        sheet_names = ", ".join(
            sheet.get("name", "")
            for sheet in file_info.get("sheets", [])
            if sheet.get("status") == "parsed"
        )
        values = [
            file_info.get("name", ""),
            file_info.get("sha256", ""),
            file_info.get("raw_row_count", 0),
            file_info.get("parsed_row_count", 0),
            sheet_names,
        ]
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_number, column, value)
        _style_body_row(worksheet, row_number, 5, styles)
        worksheet.cell(row_number, 1).alignment = styles["wrap"]
        row_number += 1

    section("警告和错误")
    _style_header(
        worksheet,
        row_number,
        ["级别", "代码", "数量", "说明", "样例"],
        styles,
    )
    row_number += 1
    all_issues = [*validation.errors, *validation.warnings]
    if not all_issues:
        all_issues = []
        values = ["信息", "NONE", 0, "没有错误或警告", ""]
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_number, column, value)
        _style_body_row(worksheet, row_number, 5, styles)
        row_number += 1
    else:
        for issue in all_issues:
            values = [
                issue.severity,
                issue.code,
                issue.count,
                issue.message,
                "\n".join(issue.samples),
            ]
            for column, value in enumerate(values, start=1):
                worksheet.cell(row_number, column, value)
            fill = (
                styles["error_fill"]
                if issue.severity == "error"
                else styles["warning_fill"]
            )
            _style_body_row(
                worksheet, row_number, 5, styles, fill=fill
            )
            worksheet.cell(row_number, 4).alignment = styles["wrap"]
            worksheet.cell(row_number, 5).alignment = styles["wrap"]
            row_number += 1

    section("分析边界")
    key_value(
        "用途",
        "产品、计费方式、后付费计费项、用量、费用及趋势分析。",
    )
    key_value(
        "限制",
        "账单不包含CPU、内存等监控数据，不能据此判断实例闲置或升降配。",
    )
    key_value(
        "建议",
        "任何优化建议仅供参考，需要结合合同、定价、监控数据和业务情况确认。",
    )

    worksheet.freeze_panes = "A4"
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 64
    worksheet.column_dimensions["C"].width = 14
    worksheet.column_dimensions["D"].width = 56
    worksheet.column_dimensions["E"].width = 70
    for row in range(1, row_number + 1):
        worksheet.row_dimensions[row].height = min(
            90, max(20, worksheet.row_dimensions[row].height or 20)
        )


def _collapse_product_rows(
    product_rows: List[Dict[str, Any]], months: List[str]
) -> List[Dict[str, Any]]:
    grouped: Dict[str, Dict[str, Any]] = {}
    for row in product_rows:
        code = row["product_code"]
        if code not in grouped:
            grouped[code] = {
                "product_code": code,
                "product_name": row["product_name"],
                "monthly_amounts": {
                    month: Decimal("0") for month in months
                },
                "total_amount": Decimal("0"),
            }
        target = grouped[code]
        for month in months:
            target["monthly_amounts"][month] += Decimal(
                row["monthly_amounts"][month]
            )
        target["total_amount"] += Decimal(row["total_amount"])
    result = []
    for row in grouped.values():
        result.append(
            {
                "product_code": row["product_code"],
                "product_name": row["product_name"],
                "monthly_amounts": {
                    month: str(value)
                    for month, value in row["monthly_amounts"].items()
                },
                "total_amount": str(row["total_amount"]),
            }
        )
    return result


def _subtitle(analysis: Dict[str, Any]) -> str:
    months = analysis["months"]
    customer = analysis.get("customer_name") or "未填写客户名称"
    currency = ", ".join(analysis.get("currencies", [])) or "币种未提供"
    return f"{customer}｜{months[0]} ~ {months[-1]}｜{currency}"


def _format_amount_columns(
    worksheet, row: int, start_column: int, end_column: int
) -> None:
    for column in range(start_column, end_column + 1):
        worksheet.cell(row, column).number_format = "#,##0.00####"


def _set_widths(
    worksheet,
    fixed: Dict[int, float],
    *,
    default_start: int,
    default_width: float,
) -> None:
    from openpyxl.utils import get_column_letter

    for column, width in fixed.items():
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(default_start, worksheet.max_column + 1):
        letter = get_column_letter(column)
        if worksheet.column_dimensions[letter].width == 13.0:
            worksheet.column_dimensions[letter].width = default_width


def _excel_number(value: Decimal) -> float:
    return float(value)


def _bold(font):
    from copy import copy

    new_font = copy(font)
    new_font.bold = True
    return new_font


def verify_report(path: Path, analysis: Dict[str, Any]) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise ValueError(
                f"Sheet不符合预期: {workbook.sheetnames} != {SHEET_NAMES}"
            )
        expected_total = Decimal(analysis["totals"]["payable_amount"])
        actual_total = _find_row_last_number(
            workbook["账单分析"], "总合计"
        )
        _assert_close(actual_total, expected_total, "账单分析总合计")

        expected_postpaid = Decimal(analysis["totals"]["postpaid_amount"])
        actual_postpaid = _find_row_last_number(
            workbook["后付费分析"], "总计"
        )
        _assert_close(actual_postpaid, expected_postpaid, "后付费分析总计")

        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str) and value.startswith(
                        ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                    ):
                        raise ValueError(
                            f"Excel包含错误值: {worksheet.title}: {value}"
                        )
    finally:
        workbook.close()


def _find_row_last_number(worksheet, label: str) -> Decimal:
    for row in worksheet.iter_rows(values_only=True):
        if row and row[0] == label:
            for value in reversed(row):
                if isinstance(value, (int, float)):
                    return Decimal(str(value))
    raise ValueError(f"未找到合计行: {worksheet.title}/{label}")


def _assert_close(actual: Decimal, expected: Decimal, label: str) -> None:
    tolerance = Decimal("0.00000001")
    if abs(actual - expected) > tolerance:
        raise ValueError(f"{label}不一致: {actual} != {expected}")
